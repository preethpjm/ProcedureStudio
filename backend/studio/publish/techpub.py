"""Exports for the TechPub web tool.

* ``ipl_excel``      — the IPL as the Excel TechPub's "Generate from Excel" route reads
  (``techback/excel_utils.py`` header names: Item No, Indent, Part Number, Keyword,
  Description, Eff Code, Units Per Assy, Attach, Torque Value, Remarks…). An IPL validated
  here can be loaded into TechPub unchanged.
* ``editor_json``    — a procedure in the TechPub procedure editor format
  (``{"task": {"title", "para": [{"intent", "para"}], "subtasks": {}}}``) with the step ids
  kept in ``meta`` so edits made in TechPub can be traced back to operations.
"""
from __future__ import annotations

import io


def ipl_excel(ipl: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "IPL"
    headers = ["Figure Number", "Item No", "Item Variant", "Indent", "Part Number", "Keyword", "Description",
               "Eff Code", "Units Per Assy", "Manufacturer Code", "Attach", "Illus Ind", "Installed-On",
               "Torque Value", "Remarks", "3D Items", "Confidence", "Flags"]
    ws.append(headers)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill("solid", fgColor="DDE7F3")
    rows = {r["id"]: r for r in ipl["rows"]}
    for r in ipl["rows"]:
        if r.get("status") == "rejected":
            continue
        nom = r["nomenclature"]
        kwd, adt = (nom.split(",", 1) + [""])[:2] if "," in nom else (nom, "")
        item = r.get("item") or ""
        variant = item[-1] if item[-1:].isalpha() else ""
        tq = r.get("torque")
        installed = rows.get(r.get("attaching_for") or r.get("parent_row") or "", {}).get("item", "")
        ws.append([r.get("fig", "1"), item.rstrip(variant) if variant else item, variant, r["indent"] - 1,
                   r["part_number"], kwd.strip(), adt.strip() or nom, r.get("effectivity") or "",
                   "RF" if r["depth"] == 0 else r.get("units_per_assy"), r.get("cage") or "",
                   "1" if r.get("attaching_for") else "", "1" if r.get("occ_ids") else "0", installed,
                   f"{tq['nominal_nm']:g} N.m" + (f" ±{tq['tolerance_nm']:g}" if tq.get("tolerance_nm") else "") if tq else "",
                   "; ".join(r.get("notes", [])), len(r.get("occ_ids", [])), r.get("confidence"),
                   " | ".join(r.get("flags", []))])
    for col, w in zip("ABCDEFGHIJKLMNOPQR", (8, 8, 6, 7, 18, 18, 30, 8, 9, 10, 7, 8, 10, 16, 30, 8, 10, 60)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def editor_json(proc: dict) -> dict:
    paras = []
    for st in proc["subtasks"]:
        paras.append({"intent": 0, "para": st["title"], "meta": {"subtask": st["id"]}})
        for s in st["steps"]:
            paras.append({"intent": 1, "para": s["text"], "meta": {"step": s["id"], "ops": s["op_ids"]}})
    return {"task": {"title": proc["title"], "para": paras, "subtasks": {}}}
